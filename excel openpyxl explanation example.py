import openpyxl

wb = openpyxl.load_workbook('nazwa.xlsx')
sheet = wb.active
sheet1 = wb['Sheet1']

for cell_obj in list(sheet.columns)[1]:
    print(cell_obj.value)
    
for row_of_cell_obj in sheet['A1':'C3']:
    for cell_obj in row_of_cell_obj:
        print(cell_obj.coordinate)


from openpyxl.utils import get_column_letter, column_index_from_string
get_column_letter(27)
get_column_letter(sheet.max_column)
column_index_from_string('A')